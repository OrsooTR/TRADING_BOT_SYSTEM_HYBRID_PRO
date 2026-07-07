from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ARTIFACT_DIR = PROJECT_ROOT / "artifacts" / "daily_reports"
VAULT_LOG_PATH = PROJECT_ROOT / "Pipeline_01" / "09_LOGS" / "report_giornata_2026-04-22.md"
WORKBOOK_PATH = PROJECT_ROOT / "artifacts" / "working_backtest_general.xlsx"
SMOKE_RESULTS_PATH = PROJECT_ROOT / "artifacts" / "derivative_backtests" / "data" / "derivatives_grid_results_smoke.csv"
PDF_PATH = ARTIFACT_DIR / "report_giornata_2026-04-22.pdf"


def read_workbook_tables() -> dict[str, list[list[object]]]:
    workbook = load_workbook(WORKBOOK_PATH, data_only=True)

    summary_rows = []
    ws_summary = workbook["SUMMARY"]
    for row in ws_summary.iter_rows(min_row=5, max_row=15, values_only=True):
        if row[0] == "TF" or row[0] is None:
            continue
        summary_rows.append(list(row))

    observations = []
    for row_idx in range(24, 31):
        value = ws_summary[f"A{row_idx}"].value
        if value:
            observations.append(str(value))

    timeframe_tables: dict[str, list[list[object]]] = {}
    for name in ("M1", "M5", "M15"):
        ws = workbook[name]
        rows = []
        for row in ws.iter_rows(min_row=2, max_row=6, values_only=True):
            rows.append(list(row))
        timeframe_tables[name] = rows

    return {
        "summary_top10": summary_rows,
        "observations": observations,
        "m1_top5": timeframe_tables["M1"],
        "m5_top5": timeframe_tables["M5"],
        "m15_top5": timeframe_tables["M15"],
    }


def read_smoke_results() -> dict[str, object]:
    frame = pd.read_csv(SMOKE_RESULTS_PATH)
    best = frame.iloc[0].to_dict()
    by_tf = {}
    for timeframe in ("M1", "M5", "M15"):
        subset = frame[frame["tf"] == timeframe].head(3).copy()
        by_tf[timeframe] = subset.to_dict(orient="records")
    return {"best": best, "by_tf": by_tf}


def significant_files_today() -> list[str]:
    candidates = [
        "Pipeline_01/07_AUTOMATION/codex_instructions.md",
        "Pipeline_01/07_AUTOMATION/memory_rules.md",
        "Pipeline_01/11_MEMORY/INDEX.md",
        "Pipeline_01/11_MEMORY/PROJECT_STATE.md",
        "Pipeline_01/11_MEMORY/DECISIONS_LOG.md",
        "Pipeline_01/11_MEMORY/DEVELOPMENT_BACKLOG.md",
        "Pipeline_01/11_MEMORY/ASSISTANT_CONTEXT.md",
        "Pipeline_01/11_MEMORY/EXPERIMENT_PROTOCOL.md",
        "src/derivatives_bt/data_utils.py",
        "src/derivatives_bt/backtest.py",
        "scripts/run_derivative_grid.py",
        "artifacts/derivative_backtests/data/derivatives_grid_results_smoke.csv",
        "artifacts/derivative_backtests/reports/derivatives_report_all_tfs_smoke.md",
    ]
    return candidates


def build_markdown(workbook_data: dict[str, list[list[object]]], smoke_data: dict[str, object]) -> str:
    best_smoke = smoke_data["best"]
    lines: list[str] = []
    lines.append("# Report Giornata - 2026-04-22")
    lines.append("")
    lines.append("## Collegamenti")
    lines.append("- [[INDEX]]")
    lines.append("- [[../11_MEMORY/INDEX]]")
    lines.append("- [[../03_STRATEGIES/TESTED/01_MA_Crossover]]")
    lines.append("- [[../03_STRATEGIES/TESTED/02_Derivative_1st_2nd_Order]]")
    lines.append("- [[../../artifacts/derivative_backtests/reports/derivatives_report_all_tfs_smoke]]")
    lines.append("- [[../../artifacts/derivative_backtests/reports/derivatives_report_m15_full]]")
    lines.append("")
    lines.append("## Executive Summary")
    lines.append("- Il progetto e stato analizzato e riorganizzato come vault Obsidian con memoria operativa centrale.")
    lines.append("- E stata costruita la base tecnica per i backtest sulle derivate con parser HistData, strategia Backtrader e runner di grid search.")
    lines.append("- E stato eseguito uno smoke test sulla strategia a derivate; la grid completa e stata preparata ma rimandata a domani.")
    lines.append("- Il backtest MA Crossover presente nel workbook resta il riferimento empirico piu maturo della giornata.")
    lines.append("")
    lines.append("## Cosa E Stato Fatto Oggi")
    lines.append("- Analisi del progetto e identificazione del suo stato reale: research vault, non bot eseguibile.")
    lines.append("- Creazione della sezione `11_MEMORY` per rendere il vault riutilizzabile senza rileggere ogni volta la chat.")
    lines.append("- Aggiornamento delle istruzioni di automazione e memory rules.")
    lines.append("- Implementazione del modulo `derivatives_bt` con loader dati, strategia e metriche.")
    lines.append("- Implementazione del runner di grid search per EUR/USD su M1, M5 e M15.")
    lines.append("- Esecuzione di smoke test ridotti e salvataggio dei risultati in `artifacts/derivative_backtests`.")
    lines.append("")
    lines.append("## Backtest MA Crossover Gia Disponibile")
    lines.append(f"- Fonte: `{WORKBOOK_PATH}`")
    lines.append("- Miglior configurazione assoluta: M1, MA 50/150, RR 4, PF 1.133, DD 55%, P&L 1018.6%.")
    lines.append("- La gerarchia dei risultati conferma che M1 domina, M5 e quasi privo di edge, M15 resta debole.")
    lines.append("")
    lines.append("### Osservazioni Workbook")
    for obs in workbook_data["observations"]:
        lines.append(f"- {obs}")
    lines.append("")
    lines.append("## Stato Derivative Backtest")
    lines.append("- Strategia definita su EMA(close), derivata prima normalizzata su ATR e derivata seconda.")
    lines.append("- Dataset pronto: EUR/USD M1 HistData 2020-2025 con resampling previsto a M5 e M15.")
    lines.append("- Smoke test completato su finestra ridotta per verificare parsing, ordini, SL/TP e metriche.")
    lines.append(
        f"- Miglior setup smoke: {best_smoke['tf']} | smooth {int(best_smoke['smooth_period'])} | "
        f"d1 {int(best_smoke['d1_period'])} | d2 {int(best_smoke['d2_period'])} | RR {best_smoke['rr']:.1f} | "
        f"OOS PF {best_smoke['oos_profit_factor']:.3f} | OOS DD {best_smoke['oos_max_drawdown_pct']:.2f}% | "
        f"OOS P&L {best_smoke['oos_pnl_pct']:.2f}%."
    )
    lines.append("- La grid completa non e stata eseguita oggi; la fase full e stata rimandata a domani.")
    lines.append("")
    lines.append("## File Significativi Creati o Aggiornati")
    for item in significant_files_today():
        lines.append(f"- `{item}`")
    lines.append("")
    lines.append("## Priorita Per Domani")
    lines.append("- Eseguire la grid completa sulle derivate con multiprocessing fuori sandbox.")
    lines.append("- Popolare il workbook Excel con tutti i risultati, buoni e cattivi.")
    lines.append("- Aggiornare il vault con report completi M1, M5, M15 e all-timeframe.")
    lines.append("- Valutare se la logica d1+d2 abbia edge sufficiente per entrare nel set di concetti candidati pre-FFT.")
    lines.append("")
    return "\n".join(lines) + "\n"


def build_table(data: list[list[object]], headers: list[str], col_widths: list[float] | None = None) -> Table:
    rows = [headers]
    for row in data:
        formatted = []
        for value in row:
            if isinstance(value, float):
                formatted.append(f"{value:.3f}" if abs(value) < 10 else f"{value:.1f}")
            else:
                formatted.append("" if value is None else str(value))
        rows.append(formatted)

    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AAB7C4")),
                ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7FAFC")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F7FAFC"), colors.HexColor("#EDF2F7")]),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    return table


def add_page_number(canvas, doc) -> None:  # type: ignore[no-untyped-def]
    canvas.saveState()
    canvas.setFont("Helvetica", 9)
    canvas.setFillColor(colors.HexColor("#5B6770"))
    canvas.drawRightString(doc.pagesize[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf(workbook_data: dict[str, list[list[object]]], smoke_data: dict[str, object]) -> None:
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=landscape(A4),
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleLarge",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=26,
        textColor=colors.HexColor("#16324F"),
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    subtitle_style = ParagraphStyle(
        "Subtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#4A5568"),
        alignment=TA_CENTER,
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "Section",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#16324F"),
        spaceBefore=6,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        alignment=TA_LEFT,
        textColor=colors.HexColor("#1F2933"),
        spaceAfter=4,
    )

    story = []
    story.append(Paragraph("Report Operativo della Giornata", title_style))
    story.append(Paragraph("Trading Bot System Hybrid PRO - 22 Aprile 2026", subtitle_style))
    story.append(
        Paragraph(
            "Documento di riepilogo completo della giornata, comprensivo dello stato del progetto, "
            "del lavoro sul vault Obsidian, della preparazione del backtest sulle derivate e del backtest "
            "MA Crossover gia presente nel workbook ricevuto.",
            body_style,
        )
    )
    story.append(Spacer(1, 6))

    story.append(Paragraph("Executive Summary", section_style))
    summary_points = [
        "Il progetto e stato riclassificato come research vault strutturato, non ancora bot eseguibile.",
        "E stata aggiunta una memory layer centrale nel vault per stato progetto, decisioni, backlog e contesto assistente.",
        "E stata implementata la base tecnica del backtest sulle derivate: parser dati, strategia Backtrader e runner batch.",
        "Lo smoke test sulle derivate e stato completato; la grid completa e stata predisposta ma rinviata a domani.",
        "Il backtest piu maturo resta il MA Crossover documentato nel workbook Excel fornito.",
    ]
    for point in summary_points:
        story.append(Paragraph(f"- {point}", body_style))

    story.append(Spacer(1, 6))
    story.append(Paragraph("Lavoro Svolto Oggi", section_style))
    work_points = [
        "Analisi dell'architettura documentale del vault e valutazione del suo livello di maturita.",
        "Creazione di una sezione `11_MEMORY` con file dedicati a project state, decisions log, development backlog, assistant context ed experiment protocol.",
        "Aggiornamento delle istruzioni operative del vault per usare Obsidian come memoria di progetto.",
        "Scaffolding tecnico del modulo Python `derivatives_bt` e dello script `run_derivative_grid.py`.",
        "Validazione con smoke test della strategia basata su derivata prima e seconda.",
    ]
    for point in work_points:
        story.append(Paragraph(f"- {point}", body_style))

    story.append(PageBreak())
    story.append(Paragraph("Backtest MA Crossover - Sintesi dal Workbook", section_style))
    story.append(
        Paragraph(
            f"Fonte analizzata: {WORKBOOK_PATH}. Il workbook riporta 390 backtest su EUR/USD dal 2020 al 2025, "
            "organizzati su SUMMARY, ALL, M1, M5 e M15.",
            body_style,
        )
    )
    headers = ["TF", "Fast MA", "Slow MA", "RR", "Trades", "Win Rate %", "Profit Factor", "Max DD %", "P&L %", "Longs", "Shorts", "Avg Bars"]
    story.append(build_table(workbook_data["summary_top10"], headers))
    story.append(Spacer(1, 8))
    for obs in workbook_data["observations"]:
        story.append(Paragraph(f"- {obs}", body_style))

    story.append(Spacer(1, 8))
    story.append(Paragraph("Top 5 per Timeframe", section_style))
    story.append(Paragraph("M1", body_style))
    story.append(build_table(workbook_data["m1_top5"], headers))
    story.append(Spacer(1, 5))
    story.append(Paragraph("M5", body_style))
    story.append(build_table(workbook_data["m5_top5"], headers))
    story.append(Spacer(1, 5))
    story.append(Paragraph("M15", body_style))
    story.append(build_table(workbook_data["m15_top5"], headers))

    story.append(PageBreak())
    story.append(Paragraph("Backtest Derivate - Stato di Avanzamento", section_style))
    smoke_best = smoke_data["best"]
    derivative_points = [
        "Strategia impostata su EMA(close), derivata prima normalizzata su ATR(14) e derivata seconda come differenza della derivata prima.",
        "Input dati pronto: EUR/USD HistData M1 2020-2025, con resampling previsto a M5 e M15.",
        "La griglia completa prevista contiene 4374 combinazioni, da eseguire in train 2020-2023 e OOS 2024-2025.",
        "Lo smoke test e servito a verificare parsing, logica ordini, gestione SL/TP, metriche e ranking.",
        "La grid completa non e stata eseguita oggi; il run full viene ripreso domani.",
    ]
    for point in derivative_points:
        story.append(Paragraph(f"- {point}", body_style))

    story.append(Spacer(1, 6))
    story.append(
        Paragraph(
            "Miglior setup emerso nello smoke test: "
            f"{smoke_best['tf']} | smooth {int(smoke_best['smooth_period'])} | d1 {int(smoke_best['d1_period'])} | "
            f"d2 {int(smoke_best['d2_period'])} | d1 thr {smoke_best['d1_threshold']:.2f} | "
            f"d2 thr {smoke_best['d2_threshold']:.2f} | SL ATR {smoke_best['sl_atr_mult']:.1f} | RR {smoke_best['rr']:.1f} | "
            f"OOS PF {smoke_best['oos_profit_factor']:.3f} | OOS DD {smoke_best['oos_max_drawdown_pct']:.2f}% | "
            f"OOS P&L {smoke_best['oos_pnl_pct']:.2f}%.",
            body_style,
        )
    )

    smoke_headers = ["TF", "Smooth", "D1", "D2", "D1 Thr", "D2 Thr", "SL ATR", "RR", "OOS PF", "OOS DD %", "OOS P&L %"]
    smoke_rows = []
    for timeframe in ("M1", "M5", "M15"):
        for row in smoke_data["by_tf"][timeframe]:
            smoke_rows.append(
                [
                    row["tf"],
                    row["smooth_period"],
                    row["d1_period"],
                    row["d2_period"],
                    row["d1_threshold"],
                    row["d2_threshold"],
                    row["sl_atr_mult"],
                    row["rr"],
                    row["oos_profit_factor"],
                    row["oos_max_drawdown_pct"],
                    row["oos_pnl_pct"],
                ]
            )
    story.append(build_table(smoke_rows, smoke_headers))

    story.append(Spacer(1, 8))
    story.append(Paragraph("File Significativi della Giornata", section_style))
    for item in significant_files_today():
        story.append(Paragraph(f"- {item}", body_style))

    story.append(Spacer(1, 8))
    story.append(Paragraph("Prossimi Passi", section_style))
    next_steps = [
        "Eseguire la grid completa sulle derivate con multiprocessing.",
        "Popolare il workbook Excel con tutti i risultati del derivative backtest.",
        "Scrivere i report finali M1, M5, M15 e all-timeframe nel vault.",
        "Decidere se la strategia a derivate merita di entrare tra i concetti candidati prima della fase FFT.",
    ]
    for point in next_steps:
        story.append(Paragraph(f"- {point}", body_style))

    doc.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)


def main() -> None:
    workbook_data = read_workbook_tables()
    smoke_data = read_smoke_results()
    markdown = build_markdown(workbook_data, smoke_data)
    VAULT_LOG_PATH.write_text(markdown, encoding="utf-8")
    build_pdf(workbook_data, smoke_data)
    print(PDF_PATH)


if __name__ == "__main__":
    main()
