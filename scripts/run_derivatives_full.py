"""
Derivatives 1st/2nd Order — Full Grid Search
=============================================
Grid: smooth_period × d1_period × d2_period × d1_threshold × d2_threshold × sl_atr_mult × rr
      2 × 2 × 2 × 3 × 1 × 3 × 3 = 216 combinations per TF  |  648 total (3 TFs)

Excluded values (rationale):
  smooth_period  5     -> too noisy (EMA ~= close), 15 -> redundant between 10 and 20
  d1_period      2     -> 2-bar derivative = pure noise; 6 -> redundant between 4 and 8
  d2_threshold   0.02  -> validate pure direction (0.00) first before tuning d2 threshold
  rr             1.5   -> trend-following derivatives don't pay below 2.0 RR

Train: 2020-01-01 → 2023-12-31
OOS  : 2024-01-01 → 2025-12-31

Outputs (all in artifacts/derivative_backtests/):
  data/derivatives_grid_{tf}_full.csv         — per-TF results
  data/derivatives_grid_full.csv/.json        — combined all TFs
  DERIVATIVES_GridSearch_EURUSD_1.xlsx        — Excel workbook (5 sheets)
"""
from __future__ import annotations

import json
import os
import sys
from datetime import UTC, datetime
from math import prod
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from derivatives_bt.backtest import RunConfig, run_optimization
from derivatives_bt.data_utils import load_histdata_zips, resample_ohlcv, split_frame

# ── paths ──────────────────────────────────────────────────────────────────────
PROJECT_ROOT  = Path(__file__).resolve().parents[1]
ARTIFACT_ROOT = PROJECT_ROOT / "artifacts" / "derivative_backtests"
DATA_ROOT     = ARTIFACT_ROOT / "data"
REPORT_ROOT   = ARTIFACT_ROOT / "reports"

HISTDATA_DIR = Path(os.environ.get("HISTDATA_DIR", PROJECT_ROOT / "data" / "histdata"))
DEFAULT_ZIPS = [HISTDATA_DIR / f"HISTDATA_COM_ASCII_EURUSD_M1{y}.zip"
                for y in [2020, 2021, 2022, 2023, 2024, 2025]]

# ── parameter grid: 2×2×2×3×1×3×3 = 216 per TF ───────────────────────────────
GRID: dict[str, list] = {
    "smooth_period": [10, 20],
    "d1_period":     [4, 8],
    "d2_period":     [2, 4],
    "d1_threshold":  [0.05, 0.10, 0.15],
    "d2_threshold":  [0.00],
    "sl_atr_mult":   [1.0, 1.5, 2.0],
    "rr":            [2.0, 3.0, 4.0],
}

TIMEFRAMES  = ["M1", "M5", "M15"]
PARAM_COLS  = ["smooth_period", "d1_period", "d2_period",
               "d1_threshold", "d2_threshold", "sl_atr_mult", "rr"]

EXPORT_COLS = [
    "tf", "smooth_period", "d1_period", "d2_period",
    "d1_threshold", "d2_threshold", "sl_atr_mult", "rr",
    "train_trades", "train_win_rate", "train_profit_factor",
    "train_max_drawdown_pct", "train_pnl_pct",
    "oos_trades", "oos_win_rate", "oos_profit_factor",
    "oos_max_drawdown_pct", "oos_pnl_pct",
    "total_trades", "total_longs", "total_shorts", "avg_bars",
]

# ── Excel column definitions: (header, df_key, col_width, num_format, group) ──
COL_DEFS: list[tuple[str, str, int, str, str]] = [
    ("TF",          "tf",                       7,  "@",      "param"),
    ("Smooth",      "smooth_period",             8,  "0",      "param"),
    ("D1 Per.",     "d1_period",                 8,  "0",      "param"),
    ("D2 Per.",     "d2_period",                 8,  "0",      "param"),
    ("D1 Thr.",     "d1_threshold",              8,  "0.00",   "param"),
    ("D2 Thr.",     "d2_threshold",              8,  "0.00",   "param"),
    ("SL ATR",      "sl_atr_mult",               7,  "0.0",    "param"),
    ("RR",          "rr",                        6,  "0.0",    "param"),
    ("Trades",      "train_trades",              8,  "#,##0",  "train"),
    ("Win Rate %",  "train_win_rate",            10, "0.0",    "train"),
    ("PF",          "train_profit_factor",       8,  "0.000",  "train"),
    ("Max DD %",    "train_max_drawdown_pct",    9,  "0.00",   "train"),
    ("P&L %",       "train_pnl_pct",             9,  "0.0",    "train"),
    ("Trades",      "oos_trades",                8,  "#,##0",  "oos"),
    ("Win Rate %",  "oos_win_rate",              10, "0.0",    "oos"),
    ("PF",          "oos_profit_factor",         8,  "0.000",  "oos"),
    ("Max DD %",    "oos_max_drawdown_pct",      9,  "0.00",   "oos"),
    ("P&L %",       "oos_pnl_pct",               9,  "0.0",    "oos"),
    ("Tot. Trades", "total_trades",              10, "#,##0",  "comb"),
    ("Longs",       "total_longs",               7,  "#,##0",  "comb"),
    ("Shorts",      "total_shorts",              7,  "#,##0",  "comb"),
    ("Avg Bars",    "avg_bars",                  8,  "0.0",    "comb"),
]

GROUP_COLORS = {
    "param": {"hdr": "2E75B6", "odd": "DEEAF1", "even": "FFFFFF"},
    "train": {"hdr": "1A5276", "odd": "D6EAF8", "even": "FFFFFF"},
    "oos":   {"hdr": "1E7145", "odd": "E2EFDA", "even": "FFFFFF"},
    "comb":  {"hdr": "7B3F00", "odd": "FCE4D6", "even": "FFFFFF"},
}
TITLE_BG = "1F3864"

_THIN   = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")


# ── Excel helpers ──────────────────────────────────────────────────────────────

def _c(ws, row: int, col: int, value, *, bold=False, size=9, color="000000",
       bg: str | None = None, fmt: str | None = None, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font      = Font(name="Arial", bold=bold, size=size, color=color)
    cell.alignment = _CENTER
    if bg:
        cell.fill = PatternFill("solid", start_color=bg)
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border
    return cell


def _group_row(ws, row: int, groups: list[tuple[str, str]]) -> None:
    col = 1
    for label, grp in groups:
        span = sum(1 for *_, g in COL_DEFS if g == grp)
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row, end_column=col + span - 1)
        _c(ws, row, col, label, bold=True, size=9, color="FFFFFF",
           bg=GROUP_COLORS[grp]["hdr"])
        col += span
    ws.row_dimensions[row].height = 16


def _header_row(ws, row: int) -> None:
    for ci, (hdr, _, w, _, grp) in enumerate(COL_DEFS, 1):
        _c(ws, row, ci, hdr, bold=True, size=9, color="FFFFFF",
           bg=GROUP_COLORS[grp]["hdr"], border=_BORDER)
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[row].height = 14


def _safe(v):
    if isinstance(v, float) and (v != v or v == float("inf") or v == float("-inf")):
        return 999.999
    return v


def _data_rows(ws, df: pd.DataFrame, start_row: int) -> None:
    for ri, row_data in enumerate(df.itertuples(index=False), start_row):
        shade = "odd" if ri % 2 else "even"
        for ci, (_, key, _, fmt, grp) in enumerate(COL_DEFS, 1):
            _c(ws, ri, ci, _safe(getattr(row_data, key)),
               bg=GROUP_COLORS[grp][shade], fmt=fmt, border=_BORDER)


def _write_data_sheet(ws, df: pd.DataFrame, title: str) -> None:
    ncols = len(COL_DEFS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    _c(ws, 1, 1, title, bold=True, size=11, color="FFFFFF", bg=TITLE_BG)
    ws.row_dimensions[1].height = 20

    _group_row(ws, 2, [
        ("PARAMETERS",        "param"),
        ("TRAIN  2020–2023",  "train"),
        ("OOS  2024–2025",    "oos"),
        ("COMBINED",          "comb"),
    ])
    _header_row(ws, 3)
    ws.freeze_panes = "A4"
    _data_rows(ws, df, 4)


def _write_summary_sheet(ws, df: pd.DataFrame, total: int, ts: str) -> None:
    ncols  = len(COL_DEFS)
    n_comb = prod(len(v) for v in GRID.values())

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    _c(ws, 1, 1, "DERIVATIVES GRID SEARCH — EURUSD 2020-2025",
       bold=True, size=13, color="FFFFFF", bg=TITLE_BG)
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    info = (f"Generated: {ts}  |  Combinations: {total}  |  "
            f"Grid: {n_comb}/TF × 3 TFs  |  Split: Train 2020-2023 / OOS 2024-2025")
    _c(ws, 2, 1, info, size=9, color="595959", bg="F2F2F2")
    ws.row_dimensions[2].height = 14

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    _c(ws, 3, 1, "TOP 15 — MIGLIORI OOS PROFIT FACTOR",
       bold=True, size=11, bg="FFFFFF")
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 6

    _group_row(ws, 5, [
        ("PARAMETERS",        "param"),
        ("TRAIN  2020–2023",  "train"),
        ("OOS  2024–2025",    "oos"),
        ("COMBINED",          "comb"),
    ])
    _header_row(ws, 6)
    ws.freeze_panes = "A7"
    _data_rows(ws, df.head(15), 7)


# ── Excel export ───────────────────────────────────────────────────────────────

def export_excel(all_df: pd.DataFrame, tf_frames: dict[str, pd.DataFrame],
                 path: Path, ts: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("SUMMARY")
    _write_summary_sheet(ws, all_df, len(all_df), ts)

    ws = wb.create_sheet("ALL")
    _write_data_sheet(ws, all_df, "DERIVATIVES — ALL TIMEFRAMES — EURUSD 2020-2025")

    for tf in TIMEFRAMES:
        ws = wb.create_sheet(tf)
        _write_data_sheet(ws, tf_frames[tf], f"DERIVATIVES — {tf} — EURUSD 2020-2025")

    wb.save(path)
    print(f"[Excel] Saved → {path}")


# ── results helpers ────────────────────────────────────────────────────────────

def _prefix(frame: pd.DataFrame, pfx: str) -> pd.DataFrame:
    return frame.rename(columns={
        "trades":           f"{pfx}_trades",
        "win_rate":         f"{pfx}_win_rate",
        "profit_factor":    f"{pfx}_profit_factor",
        "max_drawdown_pct": f"{pfx}_max_drawdown_pct",
        "pnl_pct":          f"{pfx}_pnl_pct",
        "longs":            f"{pfx}_longs",
        "shorts":           f"{pfx}_shorts",
        "avg_bars":         f"{pfx}_avg_bars",
    })


def build_combined(tf: str, train_res: list, oos_res: list) -> pd.DataFrame:
    train_df = _prefix(pd.DataFrame(train_res), "train")
    oos_df   = _prefix(pd.DataFrame(oos_res),   "oos")
    merged   = train_df.merge(oos_df, on=PARAM_COLS)
    merged.insert(0, "tf", tf)

    merged["total_trades"] = merged["train_trades"] + merged["oos_trades"]
    merged["total_longs"]  = merged["train_longs"]  + merged["oos_longs"]
    merged["total_shorts"] = merged["train_shorts"] + merged["oos_shorts"]

    tot = merged["train_trades"] + merged["oos_trades"]
    merged["avg_bars"] = (
        merged["train_avg_bars"] * merged["train_trades"]
        + merged["oos_avg_bars"] * merged["oos_trades"]
    ) / tot.replace(0, float("nan"))
    merged["avg_bars"] = merged["avg_bars"].fillna(0.0)

    merged = merged.drop(columns=[
        "train_longs", "train_shorts", "train_avg_bars",
        "oos_longs",   "oos_shorts",   "oos_avg_bars",
    ])
    return (
        merged
        .sort_values(
            ["oos_profit_factor", "oos_max_drawdown_pct", "oos_pnl_pct"],
            ascending=[False, True, False],
        )
        .reset_index(drop=True)
        [EXPORT_COLS]
    )


# ── main ───────────────────────────────────────────────────────────────────────

def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ts = datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")
    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    REPORT_ROOT.mkdir(parents=True, exist_ok=True)

    n_combos = prod(len(v) for v in GRID.values())
    total_runs = n_combos * len(TIMEFRAMES) * 2
    print("=" * 60)
    print(f"DERIVATIVES FULL GRID SEARCH — EURUSD 2020-2025")
    print(f"Grid: {n_combos} combos/TF × {len(TIMEFRAMES)} TFs × 2 periods = {total_runs} runs")
    print(f"Parameters:")
    for k, v in GRID.items():
        print(f"  {k:<16}: {v}")
    print(f"Timestamp: {ts}")
    print("=" * 60)

    # Load M1 data (use cached CSV if available, else read zips)
    cached_csv = DATA_ROOT / "eurusd_m1_2020_2025_combined.csv"
    if cached_csv.exists():
        print(f"\nLoading M1 from cache: {cached_csv}")
        m1 = pd.read_csv(cached_csv, index_col="datetime", parse_dates=True)
    else:
        print("\nLoading M1 from zips...")
        m1 = load_histdata_zips(DEFAULT_ZIPS)
        m1.to_csv(cached_csv)
    print(f"M1 bars: {len(m1):,}  ({m1.index[0].date()} -> {m1.index[-1].date()})")

    config = RunConfig()
    tf_frames: dict[str, pd.DataFrame] = {}
    all_parts: list[pd.DataFrame] = []

    for tf in TIMEFRAMES:
        print(f"\n{'─' * 40}")
        print(f"[{tf}] Resampling...")
        data = resample_ohlcv(m1, tf)
        train_df, oos_df = split_frame(
            data,
            train_start="2020-01-01", train_end="2023-12-31 23:59:59",
            oos_start="2024-01-01",   oos_end="2025-12-31 23:59:59",
        )
        print(f"[{tf}] train={len(train_df):,} bars  oos={len(oos_df):,} bars")

        print(f"[{tf}] Optimizing TRAIN ({n_combos} combos)...")
        train_res = run_optimization(train_df, GRID, config=config)
        print(f"[{tf}] Optimizing OOS ({n_combos} combos)...")
        oos_res = run_optimization(oos_df, GRID, config=config)

        combined = build_combined(tf, train_res, oos_res)
        tf_frames[tf] = combined

        out_csv = DATA_ROOT / f"derivatives_grid_{tf.lower()}_full.csv"
        combined.to_csv(out_csv, index=False)
        print(f"[{tf}] Done — {len(combined)} rows → {out_csv.name}")

        best = combined.iloc[0]
        print(f"[{tf}] Best OOS PF={best['oos_profit_factor']:.3f}  "
              f"DD={best['oos_max_drawdown_pct']:.2f}%  "
              f"Trades={int(best['oos_trades'])}  RR={best['rr']}")

        all_parts.append(combined)

    # Combine and sort all TFs
    all_df = (
        pd.concat(all_parts, ignore_index=True)
        .sort_values(
            ["oos_profit_factor", "oos_max_drawdown_pct", "oos_pnl_pct"],
            ascending=[False, True, False],
        )
        .reset_index(drop=True)
    )

    all_df.to_csv(DATA_ROOT / "derivatives_grid_full.csv", index=False)
    all_df.to_json(DATA_ROOT / "derivatives_grid_full.json", orient="records", indent=2)

    # Excel output
    xlsx_artifact = ARTIFACT_ROOT / "DERIVATIVES_GridSearch_EURUSD_1.xlsx"
    export_excel(all_df, tf_frames, xlsx_artifact, ts)

    print("\n" + "=" * 60)
    print(f"COMPLETE — {len(all_df)} total rows across {len(TIMEFRAMES)} TFs")
    print(f"Best overall: OOS PF={all_df.iloc[0]['oos_profit_factor']:.3f}  "
          f"TF={all_df.iloc[0]['tf']}  RR={all_df.iloc[0]['rr']}  "
          f"DD={all_df.iloc[0]['oos_max_drawdown_pct']:.2f}%")
    print(f"Excel → {xlsx_artifact}")
    print("=" * 60)


if __name__ == "__main__":
    main()
